# apps/administracion/infrastructure/report_views.py
"""
Reportes exportables (PDF con ReportLab, Excel con openpyxl).
Cada vista devuelve ``HttpResponse`` con el ``Content-Type`` correcto
y el header ``Content-Disposition: attachment; filename=...``.
"""

from datetime import datetime
from io import BytesIO
from typing import Iterable, List, Sequence

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from rest_framework.views import APIView

from apps.inventario.infrastructure.models import MovimientoORM, ProductoORM
from apps.terceros.infrastructure.models import ClienteORM


# ──────────────────────────────────────────────────────────────────────
# Constantes y tipos MIME
# ──────────────────────────────────────────────────────────────────────

MIME_PDF = 'application/pdf'
MIME_XLSX = (
    'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
)


# ──────────────────────────────────────────────────────────────────────
# Helpers genéricos
# ──────────────────────────────────────────────────────────────────────

def _timestamp() -> str:
    return timezone.localtime().strftime('%Y%m%d_%H%M')


def _build_pdf_response(
    titulo: str,
    encabezados: Sequence[str],
    filas: Iterable[Sequence],
    nombre_archivo: str,
    orientacion_horizontal: bool = True,
) -> HttpResponse:
    """Construye un PDF con un título, fecha de generación y una tabla."""
    buffer = BytesIO()
    pagesize = landscape(A4) if orientacion_horizontal else A4
    doc = SimpleDocTemplate(
        buffer,
        pagesize=pagesize,
        leftMargin=15 * mm,
        rightMargin=15 * mm,
        topMargin=15 * mm,
        bottomMargin=15 * mm,
        title=titulo,
    )

    styles = getSampleStyleSheet()
    story: List = [
        Paragraph(f"<b>{titulo}</b>", styles['Title']),
        Paragraph(
            f"Generado: {timezone.localtime().strftime('%Y-%m-%d %H:%M')}",
            styles['Normal'],
        ),
        Spacer(1, 8 * mm),
    ]

    data = [list(encabezados)]
    data.extend([[str(c) if c is not None else "" for c in fila] for fila in filas])

    if len(data) == 1:
        story.append(Paragraph("<i>No hay datos disponibles.</i>", styles['Italic']))
    else:
        tabla = Table(data, repeatRows=1)
        tabla.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1F2937')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 10),
            ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 1), (-1, -1), 9),
            ('GRID', (0, 0), (-1, -1), 0.25, colors.grey),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.whitesmoke, colors.white]),
            ('LEFTPADDING', (0, 0), (-1, -1), 6),
            ('RIGHTPADDING', (0, 0), (-1, -1), 6),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        story.append(tabla)

    doc.build(story)
    buffer.seek(0)
    response = HttpResponse(buffer.getvalue(), content_type=MIME_PDF)
    response['Content-Disposition'] = f'attachment; filename="{nombre_archivo}"'
    return response


def _build_excel_response(
    titulo: str,
    encabezados: Sequence[str],
    filas: Iterable[Sequence],
    nombre_archivo: str,
    nombre_hoja: str = "Reporte",
) -> HttpResponse:
    """Construye un .xlsx con una hoja con encabezado estilizado."""
    wb = Workbook()
    ws = wb.active
    ws.title = nombre_hoja[:31] or "Reporte"

    ws.append([titulo])
    ws.merge_cells(
        start_row=1, start_column=1,
        end_row=1, end_column=len(encabezados),
    )
    titulo_cell = ws.cell(row=1, column=1)
    titulo_cell.font = Font(bold=True, size=14, color="FFFFFF")
    titulo_cell.fill = PatternFill("solid", fgColor="1F2937")
    titulo_cell.alignment = Alignment(horizontal="center", vertical="center")

    ws.append([f"Generado: {timezone.localtime().strftime('%Y-%m-%d %H:%M')}"])
    ws.append([])

    ws.append(list(encabezados))
    fila_encabezados = ws.max_row
    for col_idx in range(1, len(encabezados) + 1):
        celda = ws.cell(row=fila_encabezados, column=col_idx)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="374151")
        celda.alignment = Alignment(horizontal="center", vertical="center")

    for fila in filas:
        ws.append(list(fila))

    for col_idx, _ in enumerate(encabezados, start=1):
        max_len = len(str(encabezados[col_idx - 1]))
        for row_idx in range(fila_encabezados + 1, ws.max_row + 1):
            valor = ws.cell(row=row_idx, column=col_idx).value
            if valor is not None:
                max_len = max(max_len, len(str(valor)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 60)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    response = HttpResponse(buffer.getvalue(), content_type=MIME_XLSX)
    response['Content-Disposition'] = f'attachment; filename="{nombre_archivo}"'
    return response


# ──────────────────────────────────────────────────────────────────────
# Datasets
# ──────────────────────────────────────────────────────────────────────

INVENTARIO_HEADERS = [
    "Código", "Nombre", "Categoría", "Stock actual",
    "Stock mínimo", "Ubicación", "Precio venta",
]


def _inventario_filas() -> Iterable[Sequence]:
    qs = (
        ProductoORM.objects
        .select_related('categoria', 'ubicacion')
        .filter(activo=True)
        .order_by('codigo')
    )
    for p in qs:
        yield [
            p.codigo,
            p.nombre,
            p.categoria.nombre if p.categoria_id else "",
            p.stock_actual,
            p.stock_minimo,
            p.ubicacion.nombre if p.ubicacion_id else "",
            f"{p.precio_venta:.2f}",
        ]


MOVIMIENTOS_HEADERS = [
    "Fecha", "Tipo", "Producto", "Cantidad",
    "Stock anterior", "Stock nuevo", "Usuario", "Tercero",
]


def _movimientos_filas() -> Iterable[Sequence]:
    qs = (
        MovimientoORM.objects
        .select_related(
            'producto', 'tipo_movimiento', 'usuario',
            'proveedor', 'cliente',
        )
        .order_by('-fecha')
    )
    for m in qs:
        if m.proveedor_id:
            tercero = f"Proveedor: {m.proveedor.nombre}"
        elif m.cliente_id:
            tercero = f"Cliente: {m.cliente.nombre}"
        else:
            tercero = ""
        yield [
            timezone.localtime(m.fecha).strftime('%Y-%m-%d %H:%M') if m.fecha else "",
            m.tipo_movimiento.nombre if m.tipo_movimiento_id else "",
            f"{m.producto.codigo} - {m.producto.nombre}" if m.producto_id else "",
            m.cantidad,
            m.stock_anterior,
            m.stock_nuevo,
            m.usuario.nombre if m.usuario_id and getattr(m.usuario, 'nombre', None) else "",
            tercero,
        ]


CLIENTES_HEADERS = [
    "Nombre", "Identificación", "Email", "Teléfono",
    "Dirección", "Puntos", "Nivel fidelidad", "Activo",
]


def _clientes_filas() -> Iterable[Sequence]:
    qs = ClienteORM.objects.all().order_by('nombre')
    for c in qs:
        yield [
            c.nombre,
            c.identificacion or "",
            c.email or "",
            c.telefono or "",
            c.direccion or "",
            c.puntos_fidelizacion,
            c.nivel_fidelidad,
            "Sí" if c.activo else "No",
        ]


# ──────────────────────────────────────────────────────────────────────
# Vistas — Inventario
# ──────────────────────────────────────────────────────────────────────

class ReporteInventarioPDFView(APIView):
    """GET /reportes/inventario/pdf — Reporte PDF de productos activos."""

    def get(self, request):
        return _build_pdf_response(
            titulo="Reporte de inventario — Warehouse_iq",
            encabezados=INVENTARIO_HEADERS,
            filas=_inventario_filas(),
            nombre_archivo=f"inventario_{_timestamp()}.pdf",
        )


class ReporteInventarioExcelView(APIView):
    """GET /reportes/inventario/excel — Reporte Excel de productos activos."""

    def get(self, request):
        return _build_excel_response(
            titulo="Reporte de inventario — Warehouse_iq",
            encabezados=INVENTARIO_HEADERS,
            filas=_inventario_filas(),
            nombre_archivo=f"inventario_{_timestamp()}.xlsx",
            nombre_hoja="Inventario",
        )


# ──────────────────────────────────────────────────────────────────────
# Vistas — Movimientos
# ──────────────────────────────────────────────────────────────────────

class ReporteMovimientosPDFView(APIView):
    """GET /reportes/movimientos/pdf — Histórico de movimientos."""

    def get(self, request):
        return _build_pdf_response(
            titulo="Reporte de movimientos — Warehouse_iq",
            encabezados=MOVIMIENTOS_HEADERS,
            filas=_movimientos_filas(),
            nombre_archivo=f"movimientos_{_timestamp()}.pdf",
        )


class ReporteMovimientosExcelView(APIView):
    """GET /reportes/movimientos/excel — Histórico de movimientos."""

    def get(self, request):
        return _build_excel_response(
            titulo="Reporte de movimientos — Warehouse_iq",
            encabezados=MOVIMIENTOS_HEADERS,
            filas=_movimientos_filas(),
            nombre_archivo=f"movimientos_{_timestamp()}.xlsx",
            nombre_hoja="Movimientos",
        )


# ──────────────────────────────────────────────────────────────────────
# Vistas — Clientes
# ──────────────────────────────────────────────────────────────────────

class ReporteClientesPDFView(APIView):
    """GET /reportes/clientes/pdf — Listado completo de clientes."""

    def get(self, request):
        return _build_pdf_response(
            titulo="Reporte de clientes — Warehouse_iq",
            encabezados=CLIENTES_HEADERS,
            filas=_clientes_filas(),
            nombre_archivo=f"clientes_{_timestamp()}.pdf",
        )


class ReporteClientesExcelView(APIView):
    """GET /reportes/clientes/excel — Listado completo de clientes."""

    def get(self, request):
        return _build_excel_response(
            titulo="Reporte de clientes — Warehouse_iq",
            encabezados=CLIENTES_HEADERS,
            filas=_clientes_filas(),
            nombre_archivo=f"clientes_{_timestamp()}.xlsx",
            nombre_hoja="Clientes",
        )
