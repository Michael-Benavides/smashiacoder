# apps/administracion/infrastructure/report_generators.py
"""Generadores PDF (ReportLab) y Excel (openpyxl) para informes gerenciales."""

from decimal import Decimal
from io import BytesIO
from typing import Iterable, List, Optional, Sequence

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from .report_data import SYSTEM_NAME

MIME_PDF = 'application/pdf'
MIME_XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'

HEADER_BG = colors.HexColor('#1F2937')
ALT_ROW_A = colors.whitesmoke
ALT_ROW_B = colors.white
COVER_BG = colors.HexColor('#18181B')

EXCEL_HEADER_FILL = PatternFill('solid', fgColor='1F2937')
EXCEL_HEADER_FONT = Font(bold=True, color='FFFFFF')
EXCEL_TITLE_FILL = PatternFill('solid', fgColor='18181B')
EXCEL_TITLE_FONT = Font(bold=True, size=14, color='FFFFFF')
EXCEL_SUBTOTAL_FONT = Font(bold=True)
RED_FILL = PatternFill('solid', fgColor='FFCCCC')
YELLOW_FILL = PatternFill('solid', fgColor='FFFFCC')
CURRENCY_FMT = '"$"#,##0.00'


def _timestamp() -> str:
    return timezone.localtime().strftime('%Y%m%d_%H%M')


def _gen_datetime() -> str:
    return timezone.localtime().strftime('%Y-%m-%d %H:%M')


def _fmt_money(value) -> str:
    if isinstance(value, Decimal):
        return f"${value:,.2f}"
    try:
        return f"${float(value):,.2f}"
    except (TypeError, ValueError):
        return str(value)


def _cover_callbacks(titulo: str, periodo: str, sistema: str = SYSTEM_NAME):
    gen = _gen_datetime()

    def on_first_page(canvas, doc):
        canvas.saveState()
        w, h = doc.pagesize
        canvas.setFillColor(COVER_BG)
        canvas.rect(0, 0, w, h, fill=1, stroke=0)
        canvas.setFillColor(colors.HexColor('#27272A'))
        canvas.rect(w / 2 - 35, h * 0.62, 70, 70, fill=1, stroke=0)
        canvas.setFillColor(colors.white)
        canvas.setFont('Helvetica-Bold', 22)
        canvas.drawCentredString(w / 2, h * 0.66, 'WI')
        canvas.setFont('Helvetica-Bold', 24)
        canvas.drawCentredString(w / 2, h * 0.50, titulo)
        canvas.setFont('Helvetica', 13)
        canvas.drawCentredString(w / 2, h * 0.42, sistema)
        canvas.drawCentredString(w / 2, h * 0.36, f'Período: {periodo}')
        canvas.drawCentredString(w / 2, h * 0.30, f'Generado: {gen}')
        canvas.restoreState()

    def on_later_pages(canvas, doc):
        canvas.saveState()
        w, h = doc.pagesize
        canvas.setFont('Helvetica', 8)
        canvas.setFillColor(colors.grey)
        canvas.drawString(15 * mm, 10 * mm, f'Warehouse IQ — {gen}')
        canvas.drawRightString(w - 15 * mm, 10 * mm, f'Página {doc.page}')
        canvas.restoreState()

    return on_first_page, on_later_pages


def _pdf_table(headers: Sequence[str], rows: Iterable[Sequence], bold_last: bool = False) -> Table:
    data = [list(headers)]
    row_list = list(rows)
    for row in row_list:
        data.append([str(c) if c is not None else '' for c in row])

    if not row_list:
        data.append(['Sin datos disponibles'] + [''] * (len(headers) - 1))

    style_cmds = [
        ('BACKGROUND', (0, 0), (-1, 0), HEADER_BG),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 9),
        ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
        ('FONTSIZE', (0, 1), (-1, -1), 8),
        ('GRID', (0, 0), (-1, -1), 0.25, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [ALT_ROW_A, ALT_ROW_B]),
        ('LEFTPADDING', (0, 0), (-1, -1), 5),
        ('RIGHTPADDING', (0, 0), (-1, -1), 5),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]
    if bold_last and len(data) > 1:
        style_cmds.append(('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'))

    tabla = Table(data, repeatRows=1)
    tabla.setStyle(TableStyle(style_cmds))
    return tabla


def _build_professional_pdf(
    titulo: str,
    periodo: str,
    sections: List[dict],
    nombre_archivo: str,
    landscape_mode: bool = True,
) -> HttpResponse:
    buffer = BytesIO()
    pagesize = landscape(A4) if landscape_mode else A4
    on_first, on_later = _cover_callbacks(titulo, periodo)

    doc = SimpleDocTemplate(
        buffer,
        pagesize=pagesize,
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=15 * mm,
        bottomMargin=18 * mm,
        title=titulo,
    )

    styles = getSampleStyleSheet()
    h_style = ParagraphStyle('SectionH', parent=styles['Heading2'], fontSize=12, spaceAfter=6, spaceBefore=10)
    n_style = styles['Normal']

    story: List = [PageBreak()]

    for section in sections:
        story.append(Paragraph(f"<b>{section['heading']}</b>", h_style))
        if section.get('summary'):
            for line in section['summary']:
                story.append(Paragraph(line, n_style))
            story.append(Spacer(1, 4 * mm))
        if section.get('text'):
            story.append(Paragraph(f"<pre>{section['text']}</pre>", n_style))
            story.append(Spacer(1, 4 * mm))
        if section.get('headers') and section.get('rows') is not None:
            story.append(_pdf_table(
                section['headers'],
                section['rows'],
                bold_last=section.get('bold_last', False),
            ))
            story.append(Spacer(1, 6 * mm))

    doc.build(story, onFirstPage=on_first, onLaterPages=on_later)
    buffer.seek(0)
    response = HttpResponse(buffer.getvalue(), content_type=MIME_PDF)
    response['Content-Disposition'] = f'attachment; filename="{nombre_archivo}"'
    return response


def _excel_style_header_row(ws, row_num: int, col_count: int):
    for col in range(1, col_count + 1):
        cell = ws.cell(row=row_num, column=col)
        cell.font = EXCEL_HEADER_FONT
        cell.fill = EXCEL_HEADER_FILL
        cell.alignment = Alignment(horizontal='center', vertical='center')


def _excel_title_row(ws, title: str, col_count: int):
    ws.append([title])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=col_count)
    c = ws.cell(row=1, column=1)
    c.font = EXCEL_TITLE_FONT
    c.fill = EXCEL_TITLE_FILL
    c.alignment = Alignment(horizontal='center', vertical='center')
    ws.append([f'Generado: {_gen_datetime()}'])
    ws.append([])


def _excel_autofit(ws, min_col=1, max_col=None):
    max_col = max_col or ws.max_column
    for col_idx in range(min_col, max_col + 1):
        max_len = 10
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 3, 50)


def _excel_add_summary_sheet(wb, titulo: str, items: List[tuple]):
    ws = wb.active
    ws.title = 'Resumen'
    _excel_title_row(ws, titulo, 2)
    ws.append(['Indicador', 'Valor'])
    _excel_style_header_row(ws, ws.max_row, 2)
    for label, value in items:
        ws.append([label, value])
    _excel_autofit(ws)


def _excel_add_data_sheet(wb, name: str, headers: Sequence, rows: Iterable, money_cols=None, estado_col=None):
    ws = wb.create_sheet(name[:31])
    ws.append(list(headers))
    header_row = ws.max_row
    _excel_style_header_row(ws, header_row, len(headers))
    money_cols = money_cols or []
    estado_col = estado_col  # 1-based index

    first_data = header_row + 1
    for row in rows:
        ws.append(list(row))

    last_row = ws.max_row
    for r in range(first_data, last_row + 1):
        for c in money_cols:
            cell = ws.cell(row=r, column=c)
            if isinstance(cell.value, (int, float, Decimal)):
                cell.number_format = CURRENCY_FMT

    if estado_col and last_row >= first_data:
        col_letter = get_column_letter(estado_col)
        rng = f'{col_letter}{first_data}:{col_letter}{last_row}'
        ws.conditional_formatting.add(
            rng,
            CellIsRule(operator='equal', formula=['"CRÍTICO"'], fill=RED_FILL),
        )
        ws.conditional_formatting.add(
            rng,
            CellIsRule(operator='equal', formula=['"AGOTADO"'], fill=RED_FILL),
        )
        ws.conditional_formatting.add(
            rng,
            CellIsRule(operator='equal', formula=['"BAJO"'], fill=YELLOW_FILL),
        )

    _excel_autofit(ws)
    return ws


def _ascii_chart(por_dia: dict) -> str:
    if not por_dia:
        return 'Sin movimientos en el período.'
    lines = ['Día          Entradas (u.)    Salidas (u.)', '─' * 48]
    for dia, vals in sorted(por_dia.items()):
        e_bar = '█' * min(vals.get('entrada', 0), 20)
        s_bar = '█' * min(vals.get('salida', 0), 20)
        lines.append(
            f"{dia}  E:{e_bar:<20} {vals.get('entrada', 0):>4}  "
            f"S:{s_bar:<20} {vals.get('salida', 0):>4}"
        )
    return '\n'.join(lines)


# ── PDF builders per report ───────────────────────────────────────────

def pdf_general(data) -> HttpResponse:
    r = data['resumen_ejecutivo']
    sections = [
        {
            'heading': 'Resumen Ejecutivo',
            'summary': [
                f"• Total productos activos: <b>{r['total_productos']}</b>",
                f"• Valor total del inventario: <b>{_fmt_money(r['valor_total_inventario'])}</b>",
                f"• Productos con stock bajo/crítico: <b>{r['productos_bajo_stock']}</b>",
                f"• Movimientos del mes: <b>{r['movimientos_mes']}</b>",
            ],
        },
        {
            'heading': 'Stock Crítico',
            'headers': ['Código', 'Nombre', 'Categoría', 'Stock', 'Mínimo', 'Déficit', 'Ubicación'],
            'rows': [
                [c['codigo'], c['nombre'], c['categoria'], c['stock_actual'],
                 c['stock_minimo'], c['deficit'], c['ubicacion']]
                for c in data['stock_critico']
            ],
        },
        {
            'heading': 'Distribución por Categoría',
            'headers': ['Categoría', 'Productos', 'Valor Total'],
            'rows': [
                [d['categoria'], d['cantidad'], _fmt_money(d['valor'])]
                for d in data['distribucion_categoria']
            ],
        },
        {
            'heading': 'Top 10 — Más Movimientos del Mes',
            'headers': ['Código', 'Producto', 'Movimientos'],
            'rows': [
                [t['producto__codigo'], t['producto__nombre'], t['total']]
                for t in data['top_movimientos']
            ],
        },
        {
            'heading': 'Alertas Activas Pendientes',
            'headers': ['Producto', 'Tipo', 'Mensaje', 'Fecha'],
            'rows': [
                [
                    a.producto.codigo if a.producto_id else '',
                    a.tipo,
                    (a.mensaje[:80] + '…') if len(a.mensaje) > 80 else a.mensaje,
                    timezone.localtime(a.created_at).strftime('%Y-%m-%d %H:%M'),
                ]
                for a in data['alertas']
            ],
        },
    ]
    return _build_professional_pdf(
        data['titulo'], data['periodo'], sections, f"general_{_timestamp()}.pdf",
    )


def pdf_inventario(data) -> HttpResponse:
    r = data['resumen']
    sections = [
        {
            'heading': 'Resumen',
            'summary': [
                f"• Productos activos: <b>{r['total_productos']}</b>",
                f"• Valor total inventario: <b>{_fmt_money(r['valor_total_inventario'])}</b>",
                f"• Bajo stock mínimo: <b>{r['productos_bajo_stock']}</b>",
            ],
        },
        {
            'heading': 'Detalle de Productos',
            'headers': [
                'Código', 'Nombre', 'Categoría', 'Ubicación', 'Stock', 'Mín.',
                'P. Compra', 'P. Venta', 'Valor Inv.', 'Estado',
            ],
            'rows': [
                [
                    f['codigo'], f['nombre'], f['categoria'], f['ubicacion'],
                    f['stock_actual'], f['stock_minimo'],
                    _fmt_money(f['precio_compra']), _fmt_money(f['precio_venta']),
                    _fmt_money(f['valor_inventario']), f['estado'],
                ]
                for f in data['filas']
            ],
        },
        {
            'heading': 'Subtotales por Categoría',
            'headers': ['Categoría', 'Productos', 'Valor Total'],
            'rows': [
                [s['categoria'], s['cantidad'], _fmt_money(s['valor'])]
                for s in data['subtotales_categoria']
            ],
            'bold_last': False,
        },
    ]
    return _build_professional_pdf(
        data['titulo'], data['periodo'], sections, f"inventario_{_timestamp()}.pdf",
    )


def pdf_movimientos(data) -> HttpResponse:
    r = data['resumen']
    sections = [
        {
            'heading': 'Resumen del Período',
            'summary': [
                f"• Total entradas: <b>{r['total_entradas']}</b>",
                f"• Total salidas: <b>{r['total_salidas']}</b>",
                f"• Total traslados: <b>{r['total_traslados']}</b>",
                f"• Producto más movido: <b>{r['producto_mas_movido']}</b>",
            ],
        },
        {
            'heading': 'Entradas vs Salidas por Día (ASCII)',
            'text': _ascii_chart(data['por_dia']),
        },
        {
            'heading': 'Detalle de Movimientos',
            'headers': [
                'Fecha', 'Producto', 'Tipo', 'Cant.', 'Stock Ant.', 'Stock Nuevo',
                'Usuario', 'Observaciones',
            ],
            'rows': [
                [
                    f['fecha'], f['producto'], f['tipo'], f['cantidad'],
                    f['stock_anterior'], f['stock_nuevo'], f['usuario'],
                    (f['observaciones'][:40] + '…') if len(f['observaciones']) > 40 else f['observaciones'],
                ]
                for f in data['filas']
            ],
        },
    ]
    return _build_professional_pdf(
        data['titulo'], data['periodo'], sections, f"movimientos_{_timestamp()}.pdf",
    )


def pdf_clientes(data) -> HttpResponse:
    r = data['resumen']
    nivel_txt = ' | '.join(f"{k}: {v}" for k, v in r['distribucion_nivel'].items())
    sections = [
        {
            'heading': 'Resumen',
            'summary': [
                f"• Total clientes activos: <b>{r['total_clientes']}</b>",
                f"• Distribución por nivel: <b>{nivel_txt}</b>",
                f"• Puntos en circulación: <b>{r['puntos_circulacion']:,}</b>",
            ],
        },
        {
            'heading': 'Listado de Clientes',
            'headers': ['Nombre', 'Identificación', 'Nivel', 'Puntos', 'Email', 'Teléfono'],
            'rows': [
                [f['nombre'], f['identificacion'], f['nivel'], f['puntos'], f['email'], f['telefono']]
                for f in data['filas']
            ],
        },
        {
            'heading': 'Top 10 — Mayor Puntaje',
            'headers': ['Nombre', 'Nivel', 'Puntos'],
            'rows': [
                [t['nombre'], t['nivel'], t['puntos']] for t in data['top_puntos']
            ],
        },
        {
            'heading': 'Canjes Recientes',
            'headers': ['Fecha', 'Cliente', 'Puntos', 'Recompensa', 'Usuario'],
            'rows': [
                [c['fecha'], c['cliente'], c['puntos'], c['recompensa'], c['usuario']]
                for c in data['canjes_recientes']
            ],
        },
    ]
    return _build_professional_pdf(
        data['titulo'], data['periodo'], sections, f"clientes_{_timestamp()}.pdf",
    )


def pdf_proveedores(data) -> HttpResponse:
    sections = [
        {
            'heading': 'Resumen',
            'summary': [
                f"• Proveedores activos: <b>{data['total_proveedores']}</b>",
            ],
        },
    ]
    for prov in data['proveedores']:
        sections.append({
            'heading': f"Proveedor: {prov['nombre']}",
            'summary': [
                f"RUC/NIT: {prov['ruc_nit'] or '—'} | Email: {prov['email'] or '—'} | "
                f"Tel: {prov['telefono'] or '—'} | Total unidades recibidas: <b>{prov['total_unidades']}</b>",
            ],
            'headers': ['Código', 'Producto', 'Unidades Recibidas'],
            'rows': [
                [p['codigo'], p['nombre'], p['unidades']]
                for p in prov['productos']
            ] or [['—', 'Sin entradas registradas', 0]],
        })
    return _build_professional_pdf(
        data['titulo'], data['periodo'], sections, f"proveedores_{_timestamp()}.pdf",
    )


# ── Excel builders ────────────────────────────────────────────────────

def _excel_response(wb: Workbook, nombre: str) -> HttpResponse:
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    response = HttpResponse(buffer.getvalue(), content_type=MIME_XLSX)
    response['Content-Disposition'] = f'attachment; filename="{nombre}"'
    return response


def excel_general(data) -> HttpResponse:
    wb = Workbook()
    r = data['resumen_ejecutivo']
    _excel_add_summary_sheet(wb, data['titulo'], [
        ('Período', data['periodo']),
        ('Total productos', r['total_productos']),
        ('Valor total inventario', float(r['valor_total_inventario'])),
        ('Productos stock bajo/crítico', r['productos_bajo_stock']),
        ('Movimientos del mes', r['movimientos_mes']),
    ])
    ws = wb['Resumen']
    ws['B4'].number_format = CURRENCY_FMT

    _excel_add_data_sheet(
        wb, 'Stock Crítico',
        ['Código', 'Nombre', 'Categoría', 'Stock', 'Mínimo', 'Déficit', 'Ubicación'],
        [[c['codigo'], c['nombre'], c['categoria'], c['stock_actual'],
          c['stock_minimo'], c['deficit'], c['ubicacion']] for c in data['stock_critico']],
    )
    _excel_add_data_sheet(
        wb, 'Por Categoría',
        ['Categoría', 'Productos', 'Valor Total'],
        [[d['categoria'], d['cantidad'], float(d['valor'])] for d in data['distribucion_categoria']],
        money_cols=[3],
    )
    _excel_add_data_sheet(
        wb, 'Top Movimientos',
        ['Código', 'Producto', 'Movimientos'],
        [[t['producto__codigo'], t['producto__nombre'], t['total']] for t in data['top_movimientos']],
    )
    _excel_add_data_sheet(
        wb, 'Alertas',
        ['Producto', 'Tipo', 'Mensaje', 'Fecha'],
        [
            [a.producto.codigo if a.producto_id else '', a.tipo, a.mensaje,
             timezone.localtime(a.created_at).strftime('%Y-%m-%d %H:%M')]
            for a in data['alertas']
        ],
    )
    return _excel_response(wb, f"general_{_timestamp()}.xlsx")


def excel_inventario(data) -> HttpResponse:
    wb = Workbook()
    r = data['resumen']
    _excel_add_summary_sheet(wb, data['titulo'], [
        ('Período', data['periodo']),
        ('Productos activos', r['total_productos']),
        ('Valor total inventario', float(r['valor_total_inventario'])),
        ('Bajo stock mínimo', r['productos_bajo_stock']),
    ])
    wb['Resumen']['B4'].number_format = CURRENCY_FMT

    _excel_add_data_sheet(
        wb, 'Inventario',
        ['Código', 'Nombre', 'Categoría', 'Ubicación', 'Stock', 'Mínimo',
         'Precio Compra', 'Precio Venta', 'Valor Inventario', 'Estado'],
        [
            [f['codigo'], f['nombre'], f['categoria'], f['ubicacion'],
             f['stock_actual'], f['stock_minimo'],
             float(f['precio_compra']), float(f['precio_venta']),
             float(f['valor_inventario']), f['estado']]
            for f in data['filas']
        ],
        money_cols=[7, 8, 9],
        estado_col=10,
    )
    _excel_add_data_sheet(
        wb, 'Subtotales',
        ['Categoría', 'Productos', 'Valor Total'],
        [[s['categoria'], s['cantidad'], float(s['valor'])] for s in data['subtotales_categoria']],
        money_cols=[3],
    )
    return _excel_response(wb, f"inventario_{_timestamp()}.xlsx")


def excel_movimientos(data) -> HttpResponse:
    wb = Workbook()
    r = data['resumen']
    _excel_add_summary_sheet(wb, data['titulo'], [
        ('Período', data['periodo']),
        ('Total entradas', r['total_entradas']),
        ('Total salidas', r['total_salidas']),
        ('Total traslados', r['total_traslados']),
        ('Producto más movido', r['producto_mas_movido']),
    ])
    _excel_add_data_sheet(
        wb, 'Por Día',
        ['Día', 'Entradas (u.)', 'Salidas (u.)', 'Traslados (u.)'],
        [
            [str(dia), v.get('entrada', 0), v.get('salida', 0), v.get('traslado', 0)]
            for dia, v in sorted(data['por_dia'].items())
        ],
    )
    _excel_add_data_sheet(
        wb, 'Movimientos',
        ['Fecha', 'Producto', 'Tipo', 'Cantidad', 'Stock Ant.', 'Stock Nuevo', 'Usuario', 'Observaciones'],
        [
            [f['fecha'], f['producto'], f['tipo'], f['cantidad'],
             f['stock_anterior'], f['stock_nuevo'], f['usuario'], f['observaciones']]
            for f in data['filas']
        ],
    )
    return _excel_response(wb, f"movimientos_{_timestamp()}.xlsx")


def excel_clientes(data) -> HttpResponse:
    wb = Workbook()
    r = data['resumen']
    items = [
        ('Período', data['periodo']),
        ('Total clientes', r['total_clientes']),
        ('Puntos en circulación', r['puntos_circulacion']),
    ]
    for nivel, cnt in r['distribucion_nivel'].items():
        items.append((f'Clientes {nivel}', cnt))
    _excel_add_summary_sheet(wb, data['titulo'], items)

    _excel_add_data_sheet(
        wb, 'Clientes',
        ['Nombre', 'Identificación', 'Nivel', 'Puntos', 'Email', 'Teléfono'],
        [[f['nombre'], f['identificacion'], f['nivel'], f['puntos'], f['email'], f['telefono']]
         for f in data['filas']],
    )
    _excel_add_data_sheet(
        wb, 'Top Puntos',
        ['Nombre', 'Nivel', 'Puntos'],
        [[t['nombre'], t['nivel'], t['puntos']] for t in data['top_puntos']],
    )
    _excel_add_data_sheet(
        wb, 'Canjes',
        ['Fecha', 'Cliente', 'Puntos', 'Recompensa', 'Usuario'],
        [[c['fecha'], c['cliente'], c['puntos'], c['recompensa'], c['usuario']]
         for c in data['canjes_recientes']],
    )
    return _excel_response(wb, f"clientes_{_timestamp()}.xlsx")


def excel_proveedores(data) -> HttpResponse:
    wb = Workbook()
    _excel_add_summary_sheet(wb, data['titulo'], [
        ('Período', data['periodo']),
        ('Proveedores activos', data['total_proveedores']),
    ])
    rows = []
    for prov in data['proveedores']:
        if prov['productos']:
            for p in prov['productos']:
                rows.append([prov['nombre'], prov['ruc_nit'], p['codigo'], p['nombre'], p['unidades']])
        else:
            rows.append([prov['nombre'], prov['ruc_nit'], '—', 'Sin entradas', 0])
    _excel_add_data_sheet(
        wb, 'Proveedores',
        ['Proveedor', 'RUC/NIT', 'Código Producto', 'Producto', 'Unidades Recibidas'],
        rows,
    )
    return _excel_response(wb, f"proveedores_{_timestamp()}.xlsx")


# ── Registry ──────────────────────────────────────────────────────────

PDF_BUILDERS = {
    'general': pdf_general,
    'inventario': pdf_inventario,
    'movimientos': pdf_movimientos,
    'clientes': pdf_clientes,
    'proveedores': pdf_proveedores,
}

EXCEL_BUILDERS = {
    'general': excel_general,
    'inventario': excel_inventario,
    'movimientos': excel_movimientos,
    'clientes': excel_clientes,
    'proveedores': excel_proveedores,
}

REPORTE_LABELS = {
    'general': 'Informe General de Bodega',
    'inventario': 'Informe de Inventario',
    'movimientos': 'Informe de Movimientos',
    'clientes': 'Informe de Clientes y Fidelización',
    'proveedores': 'Informe de Proveedores',
}


def generar_reporte_http(tipo: str, formato: str, fecha_desde=None, fecha_hasta=None) -> HttpResponse:
    if formato == 'pdf':
        builder = PDF_BUILDERS[tipo]
    else:
        builder = EXCEL_BUILDERS[tipo]

    if tipo == 'movimientos':
        from .report_data import datos_reporte_movimientos
        data = datos_reporte_movimientos(fecha_desde, fecha_hasta)
    else:
        from . import report_data
        fn = getattr(report_data, f'datos_reporte_{tipo}')
        data = fn()

    return builder(data)


def generar_reporte_bytes(tipo: str, formato: str, fecha_desde=None, fecha_hasta=None):
    response = generar_reporte_http(tipo, formato, fecha_desde, fecha_hasta)
    nombre = response['Content-Disposition'].split('filename=')[1].strip('"')
    return response.content, nombre, response['Content-Type']
